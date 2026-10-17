import logging
from io import BytesIO
from pathlib import Path

logger = logging.getLogger(__name__)


class TextExtractionError(Exception):
    """The message is safe to show to the user."""


def extract_text_from_bytes(data, ext):
    """Return the plain text of a PDF, DOCX, TXT or MD file given its bytes and extension."""
    ext = ext.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise TextExtractionError("Ce PDF est protégé par un mot de passe. Exportez-en une version sans protection.")
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if ext == ".docx":
            import docx

            document = docx.Document(BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        if ext in (".txt", ".md"):
            return data.decode("utf-8", errors="ignore")
    except TextExtractionError:
        raise
    except Exception as exc:
        logger.exception("Text extraction failed for a %s file", ext)  # full traceback in the runserver terminal
        raise TextExtractionError(
            f"Impossible de lire ce fichier ({type(exc).__name__}). Réexportez-le en PDF ou DOCX, ou collez le texte à la main."
        ) from exc
    raise TextExtractionError("Format non pris en charge. Utilisez PDF, DOCX ou TXT.")


def extract_text(path):
    path = Path(path)
    return extract_text_from_bytes(path.read_bytes(), path.suffix)
